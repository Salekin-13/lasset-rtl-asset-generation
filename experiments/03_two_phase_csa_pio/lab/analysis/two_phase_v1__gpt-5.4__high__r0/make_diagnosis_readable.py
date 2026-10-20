"""Readable prompt-failure workbook for the two-phase PIO run, revised after a first-time-reader review and an accuracy
review (2026-10-11).

  PROMPT_FAILURE_DIAGNOSIS.xlsx   Start here, Glossary, Failure groups, False positives, Missed entries, Correct finds,
                                  Evidence, Prompt rules cited, Run setup, Calls
  PROMPT_FAILURE_EXACT_TEXT.xlsx  the exact system prompts, user messages and model replies, one line per row

Sources: the run records, table.json (scorer verdicts, self-tested), fp_analysis_workflow.json (AI-agent open coding,
groups, judges, drafted rules) and rule_check.py (code tests of the rules). The exact texts are checked against the run
before anything is written. The previous layout is kept as make_diagnosis_readable_prev.py.bak.

  python make_diagnosis_readable.py
"""
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
LAB = HERE.parents[1]
sys.path.insert(0, str(LAB))
sys.path.insert(0, str(HERE))
import chain as CH  # noqa: E402
import check_lab_prompt as C  # noqa: E402
import cost_estimate as CE  # noqa: E402
import labkit as K  # noqa: E402
import rule_check as RC  # noqa: E402

VERSION, MODEL, EFFORT, REP = "two_phase_v1", "gpt-5.4", "high", 0
OUT = CH.run_dir(VERSION, MODEL, REP, EFFORT)
PROMPTS = LAB / "prompts" / VERSION

# reading order: prompt-rule fixes (code-checkable, then judgement only), scoring-only fixes, no fix
ORDER = ["B01", "B02", "B03", "B15", "B05", "B04", "B07", "B10", "B14", "B06", "B09", "B08", "B13", "B12", "B11"]
GID = {b: f"G{i:02d}" for i, b in enumerate(ORDER, 1)}
CODE_TEST = {"B01": "NC-B01", "B02": "NC-B02", "B03": "NC-B03", "B15": "NC-B15", "B05": "NC-B05", "B04": "NC-B04",
             "B07": "NC-B07", "B06": "E-B06", "B09": "E-B09"}
TEXT = {  # what it is / what should happen / proposed fix / fix belongs in
    "B01": ("Whole transaction record listed (bus request or response, ICC channel, fetch bundle)",
            "Not primary. A whole transaction record only carries traffic. List the element that stores or computes the value; if the unit has none, list nothing and say so.",
            "Prompt rule: a whole transaction record is secondary.", "Prompt rule (code can check it)"),
    "B02": ("Port field that only decides (strobe, opcode, select, byte enable)",
            "Not listed. Name the element inside the unit that makes the decision.",
            "Prompt rule: do not list a port field that only decides or signals.", "Prompt rule (code can check it)"),
    "B03": ("Field of a bus port that the unit forwards or issues",
            "Not primary. The field only carries the transaction across the unit.",
            "Prompt rule: the fields of a forwarded or issued bus port are secondary.", "Prompt rule (code can check it)"),
    "B15": ("Field of the bus port of a memory or peripheral that ends the bus (shared write word, read-back word, address, strobe, write flag, acknowledge)",
            "Not primary. List the registers behind the port; the shared words and the access fields are not primary.",
            "Prompt rule: the fields of such a bus port are secondary; list the registers behind it.", "Prompt rule (code can check it)"),
    "B05": ("Data wire into or out of a FIFO",
            "Not primary. List the data where it enters, is held or leaves the unit.",
            "Prompt rule: FIFO data hand-off wires are secondary. Note: in trng the FIFO write wire fifo.wdata (R371) is the same wire as the "
            "missed reference entry data_o (FN20). This rule removes it; the {B08} scoring alias would instead count it as a match.",
            "Prompt rule (code can check it)"),
    "B04": ("Next-value twin of a register that is already listed (name ends in _nxt)",
            "Not primary. The register itself is the asset and is already listed.",
            "Prompt rule: a register's next-value signal is secondary when the register is listed.", "Prompt rule (code can check it)"),
    "B07": ("Record field picked by its name",
            "Not primary. The unit only passes the whole record on. Exception: ctrl.pc_ret (R231) is also used by itself (line 330), so the rule keeps it.",
            "Prompt rule: do not single out a field the unit never uses by itself (removes 6 of the 7).", "Prompt rule (code can check it)"),
    "B10": ("Forwarding copy, delay or read-back word labelled primary",
            "Not primary. It only copies, delays or packs values; rule 2.6 already says so.",
            "Prompt rule for internal signals that only copy, delay or assemble a read-back word. Expected to remove only about 2 of the 12; "
            "the other 10 fall under exceptions that also protect correct finds.", "Prompt rule (needs judgement)"),
    "B14": ("Register that holds an input value for one operation",
            "Not primary. It holds a copy only while the value moves through.",
            "Prompt rule: such a register is secondary.", "Prompt rule (needs judgement)"),
    "B06": ("Interrupt-enable mask bit (a record field)",
            "The model followed the prompt (rule 2.6 makes a stored setting primary). This run's reference leaves these masks out, but the "
            "reference does list gpio's interrupt enables in the held-out modules.",
            "Scoring filter, limited to record fields named irq_*; on plain signals it would drop 6 held-out reference entries. "
            "No published basis for a prompt rule.", "Scoring filter"),
    "B09": ("Element of a unit that only sequences resets",
            "The reference does not cover this unit at all.",
            "Scoring filter for units whose only outputs are resets. No published basis for a prompt rule.", "Scoring filter"),
    "B08": ("Same RTL thing as a reference entry, written differently",
            "Counted as a false positive only because of naming: another field of a record the reference lists once (4), the wire on the "
            "other side of a port map (2), the renamed port (1), or a likely counterpart (1, a judgement).",
            "Scoring alias, no prompt change: R164, R165 and R196 would count as matches; R277, R282, R284 and R285 as duplicates of a "
            "credited record; R202 is left alone (a judgement).", "Scoring alias"),
    "B13": ("Setting, state or decision register the reference leaves out",
            "The reference keeps some registers of this shape and not others. No fact in the RTL separates these from correct finds.",
            "None.", "No fix"),
    "B12": ("Stored data at a point the reference does not use",
            "The reference records the same data path at another point. No fact in the RTL separates these from correct finds.",
            "None.", "No fix"),
    "B11": ("Internal status flag (fill level, busy, error)",
            "The reference keeps some flags of this exact shape and leaves out others.",
            "None. A drafted rule would remove 17 of the 21 but hide 1 correct find (R379 fifo.avail).", "No fix"),
}
TEXT["B05"] = TEXT["B05"][:2] + (TEXT["B05"][2].format(B08=GID["B08"]),) + TEXT["B05"][3:]
# missed entries whose same thing was listed under a name the code cannot link (from the B08 group reading; checked below)
JUDGED_PAIRS = {("neorv32_cache", "inval_i"): ("R196", "The model listed cache_o.cmd_inv (R196), wired to inv_i at line 284; "
                                                       "inv_i is the current name of inval_i."),
                ("neorv32_cache", "cache_o.cmd_dir"): ("R202", "The model listed ctrl.buf_dir (R202), the likely current counterpart "
                                                               "of cache_o.cmd_dir (a judgement).")}
KIND = {"top-level port": "port", "top-level port field": "port field", "signal": "signal", "signal field": "record field",
        "child-entity port": "port of a sub-unit", "child-entity port field": "port field of a sub-unit"}
WHO_FILL = {"Prompt told it to": "F8CBAD", "Partly broke a rule": "FFE699", "Model broke a rule": "FF7C80", "Reference choice": "BDD7EE",
            "Prompt rule (labelled secondary)": "F8CBAD", "Model": "FF7C80", "Reference problem": "BDD7EE", "Naming (scoring)": "BDD7EE"}
FIX_FILL = {"Prompt rule": "C6EFCE", "Scoring filter": "E4DFEC", "Scoring alias": "E4DFEC", "No fix": "D9D9D9"}
CLAUSES = [  # (name, file, start pattern, end pattern, in plain words)
    ("2.6 Primary asset", "common.md", r"^2\.6 ", r"^2\.7 ",
     "Primary = where the value is stored or computed in this unit, or where it crosses the unit's boundary whole (the input port "
     "that sets it, the output port that returns it unmixed). For a decision or a setting, the element that holds or makes it. "
     "An internal element that only passes the value on (a forwarding wire, a copy, a multiplexer path, a read-back word) is a "
     "candidate secondary."),
    ("2.7 Secondary asset", "common.md", r"^2\.7 ", r"^2\.8 ",
     "Secondary = not primary, and either carries the primary's value (direction a) or can change, control, gate or block it, "
     "or decides its release (direction b)."),
    ("2.9 Points of observation and influence", "common.md", r"^2\.9 ", r"^3\. ",
     "Observation point = where the asset can be seen; influence point = where it can be changed, gated or blocked. At a point "
     "that neither stores nor sets the asset, designed behaviour is not a problem; at a point that stores or sets it, the "
     "designed access still counts."),
    ("Black-box rule (section 1)", "common.md", r"^- An instance of another IP is a black box", r"^- Decide from",
     "For an instance of an IP that is not in the file, assume each of its inputs can reach each element it drives; one instance deep."),
    ("Phase 1 step 2 (Inside)", "phase1.md", r"^2\. Inside", r"^3\. ",
     "Name the internal state, and each enable, select, override or lock, as information of its own."),
    ("Phase 1 step 3 (Merge and split)", "phase1.md", r"^3\. Merge", r"^4\. ",
     "Merge pieces that always travel together; the same information repeated per index or per port of one role is one piece."),
    ("Phase 1 step 6 (Anchor)", "phase1.md", r"^6\. Anchor", r"^Rules:",
     "Give the ports where each concept enters and leaves (its anchor ports). A port that only decides is not listed."),
    ("Phase 2 step 1 (Walk)", "phase2.md", r"^1\. Walk", r"^2\. ",
     "Walk the diagram from the anchor ports and the blocks that hold the concept; every block reached is a point."),
    ("Phase 2 step 3 (Structural assets)", "phase2.md", r"^3\. Structural", r"^4\. ",
     "At each point with a 'yes', the elements that store, carry, produce or control the concept are structural assets; every "
     "anchor port is one too, even when every answer is 'no'."),
    ("Phase 2 step 4 (Primary)", "phase2.md", r"^4\. Primary", r"^5\. ", "Label primary by rule 2.6."),
    ("Phase 2 step 5 (Secondary)", "phase2.md", r"^5\. Secondary", r"^Rules:",
     "Label secondary by rule 2.7, checking the primary's direct partners in the closed set."),
]
GLOSSARY = [
    ("RTL", "The hardware design source (VHDL). Line numbers are those of the numbered RTL given to the model."),
    ("Module", "One RTL file of the processor, for example bus or uart. 15 modules were analysed (the tuning set)."),
    ("Entity / sub-unit", "A design unit declared in a module. One module can hold several entities; element names are unique only within an entity."),
    ("Element", "A port, signal or record field named in the closed set: the things the model may list."),
    ("Port / port field", "An input or output of an entity. A record port has fields (bus_req_i.data is a field of bus_req_i)."),
    ("Record / record field", "A signal that bundles several values; each value is a field (ctrl.enable)."),
    ("Closed set", "Input file listing every element with facts (direction or type, tags, partners, lines). The model may list only these names."),
    ("Diagram", "Input file drawing the RTL as blocks (processes, instances) and arrows (==> a value moves; --> only decides)."),
    ("Reference list", "The project's ground-truth list of primary assets per module (data/ground_truth), used by the strict scorer (src/eval_assets.py)."),
    ("Correct find (TP)", "A listed element that is in the reference list."),
    ("False positive (FP)", "A listed element that is not in the reference list."),
    ("Missed entry (FN)", "A reference entry that the model did not list."),
    ("Precision / recall", "Precision = share of listed elements that are correct. Recall = share of reference entries found."),
    ("Tuning set / held-out set", "Tuning = the 15 modules analysed here; the fixes were written from their errors, so gains measured on them are in-sample. Held-out = 26 other modules, not read while writing the fixes."),
    ("Primary / secondary asset", "Primary = the element that holds the protected value or decision (prompt rule 2.6). Secondary = an element on the path to it (rule 2.7). Only primaries are scored."),
    ("Concept (conceptual asset)", "The information worth protecting, named in Phase 1 (for example 'watchdog timeout'). Concept ids C1, C2 ... restart in every module."),
    ("Phase 1 / Phase 2", "Two calls per module. Phase 1 names the concepts; Phase 2 walks the diagram and labels elements primary or secondary."),
    ("PIO / point", "Points of influence and observation (IEEE P3164): places where a concept can be seen or changed. Here a point is a diagram block."),
    ("Anchor port", "A port where a concept enters or leaves the unit (Phase 1 step 6)."),
    ("'Same role' copy", "When several ports differ only by an index (dev_00 ... dev_31), the model writes one row and names the others as the same role; each copy is scored as its own element."),
    ("Objective (C / I / A)", "Confidentiality, integrity, availability: what could be broken for a concept."),
    ("Closed-set tags", "Facts per element: condition / selector / index (it decides an assignment), copied / operand (its value moves on), held on a clock edge, set by an instance, and so on."),
    ("Who (false positives)", "Prompt told it to: the model followed the prompt. Partly broke a rule: it followed one clause and broke another, or two clauses point different ways. Model broke a rule: the prompt says not to. Reference choice: the model followed the prompt but the reference chose otherwise."),
    ("AI readers and judges", "All labels come from AI agents, not from a human audit: 7 readers coded the false positives, one agent grouped them and a critic checked the groups, two judges read every correct find against the groups and the drafted rules, and one judge read each false positive against the rules' wording. Method: ERROR_ANALYSIS.md."),
    ("Code test", "rule_check.py: the mechanical version of each proposed rule, applied to the stored rows and to every reference entry."),
    ("Fingerprint (sha12)", "The first 12 hex characters of the SHA-256 of a text; equal fingerprints mean identical texts."),
]


def sha12(t):
    return hashlib.sha256(t.encode("utf-8")).hexdigest()[:12]


def words(t):
    return re.sub(r"\bTPs?\b", "correct find(s)", re.sub(r"\bFPs?\b", "false positive(s)", str(t or "")))


def fix_family(where):
    return next(k for k in FIX_FILL if where.startswith(k))


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    import eval_assets as ea
    t = json.loads((HERE / "table.json").read_text(encoding="utf-8"))
    wf = json.loads((HERE / "fp_analysis_workflow.json").read_text(encoding="utf-8"))
    code = {c["id"]: c for c in wf["codes"]}
    groups = {b["bucket_id"]: b for b in wf["clustering"]["buckets"]}
    assert set(groups) == set(ORDER), sorted(set(groups) ^ set(ORDER))
    assert {"R277", "R282", "R284", "R285", "R164", "R165", "R196", "R202"} <= set(groups["B08"]["member_ids"])
    group_of = {i: b for b, g in groups.items() for i in g["member_ids"]}
    rule_of = defaultdict(list)
    for r in wf["rules"]["rules"]:
        for b in r["bucket_ids"]:
            rule_of[b].append(r["rule_id"])
    over = [defaultdict(set), defaultdict(set)]
    for i, o in enumerate(wf["overlap"]):
        for x in (o or {}).get("tp_in_bucket", []):
            over[i][x["bucket_id"]].add(x["tp_id"])
    sup = [defaultdict(set), defaultdict(set)]           # two judges over every correct find
    for i, sp in enumerate(wf["suppression_tp"]):
        for j in (sp or {}).get("judgements", []):
            for rid in j["suppressed_by"]:
                sup[i][rid].add(j["id"])
    supfp = defaultdict(set)                             # one judge per false positive (two batches)
    for sp in wf["suppression_fp"]:
        for j in (sp or {}).get("judgements", []):
            for rid in j["suppressed_by"]:
                supfp[rid].add(j["id"])
    mods = K.gt_modules()["tuning"]
    with CH.at_root():
        refs = ea.load_refs()["gt"]
        score = ea.score(str(OUT / "assets"), refs, strict=True, only=set(mods))
    recs = {(m, ph): json.loads((OUT / f"phase{ph}" / f"{m}.json").read_text(encoding="utf-8")) for m in mods for ph in (1, 2)}
    p1c = {m: recs[(m, 1)]["reply"]["conceptual assets"] for m in mods}
    cname = {(m, c["id"]): c.get("concept") for m in mods for c in p1c[m]}
    model_name = recs[(mods[0], 1)]["usage"]["model"]
    total_usd = sum(CE.usd(r["usage"]) for r in recs.values())
    rows = {r["id"]: r for r in t["rows"]}
    for (m, rn), (rid, _) in JUDGED_PAIRS.items():
        assert rows[rid]["module"] == m and any(f["module"] == m and f["reference entry"] == rn for f in t["false negatives"])
    assert re.search(r"^\s*284 \|\s*inv_i\s*=>\s*cache_o\.cmd_inv", CH.inputs("neorv32_cache")["rtl.txt"], re.M)

    # ------------------------------------------------ code tests per row (rule_check), post-hoc scores, reference exposure
    cache = {}

    def F(m):
        if m not in cache:
            cache[m] = RC.facts(m)
        return cache[m]
    listed = defaultdict(set)
    for r in t["rows"]:
        listed[(r["module"], r["entity"])].add(r["element"])
    removed_by = defaultdict(set)
    for r in t["rows"]:
        f, ep = F(r["module"])
        x = f.get((r["entity"], r["element"]))
        if x is not None:
            removed_by[r["id"]] = {k for k, v in RC.tests(x, ep, listed[(r["module"], r["entity"])]).items() if v}
    assert not any(removed_by[r["id"]] for r in t["rows"] if r["verdict"] == "TP"), "a code test removes a correct find"

    def post_hoc(rule_set):
        run = defaultdict(list)
        for r in t["rows"]:
            if not removed_by[r["id"]] & set(rule_set):
                run[r["module"]].append((r["entity"], r["element"], r["objective listed"]))
        with CH.at_root():
            return ea.score(dict(run), refs, strict=True, only=set(mods))
    ph_prompt, ph_all = post_hoc(RC.PROMPT_RULES), post_hoc(RC.PROMPT_RULES + RC.EVAL_RULES)
    exposure = {}
    for split, ms in K.gt_modules().items():
        if split not in ("tuning", "heldout"):
            continue
        n = unresolved = hits = 0
        for m in ms:
            f, ep = F(m)
            for rname, _ in refs.get(m, []):
                n += 1
                cands = [f[k] for k in f if k[1] == rname]
                if not cands:
                    unresolved += 1
                    continue
                hits += any(any(RC.tests(x, ep, {nm for (e, nm) in f if e == x["entity"]}).values()) for x in cands)
        exposure[split] = (n, unresolved, hits)

    # ------------------------------------------------ exact texts, checked
    sys1, sys2 = C.system(VERSION, 1), C.system(VERSION, 2)
    exact = [("P1 system", "system prompt, phase 1", sys1, sha12(sys1)), ("P2 system", "system prompt, phase 2", sys2, sha12(sys2))]
    for m in mods:
        mm = m.replace("neorv32_", "")
        for ph in (1, 2):
            r = recs[(m, ph)]
            user = CH.user_message(VERSION, ph, m, p1c[m] if ph == 2 else None)
            assert sha12(user) == r["user sha"], f"{m} phase {ph}: rebuilt user message differs from the one sent"
            assert sha12(sys1 if ph == 1 else sys2) == r["system sha"], f"{m} phase {ph}: system prompt differs"
            exact.append((f"{mm} P{ph} user", f"{m}, phase {ph}: user message", user, r["user sha"]))
            exact.append((f"{mm} P{ph} reply", f"{m}, phase {ph}: model reply", r["raw"], sha12(r["raw"])))

    # ------------------------------------------------ evidence helpers
    rtl = {m: {int(mm.group(1)): l.strip() for l in CH.inputs(m)["rtl.txt"].splitlines() if (mm := re.match(r"\s*(\d+) \|", l))} for m in mods}
    names = {m: CH.closed_names(m) for m in mods}
    diag = {}
    for m in mods:
        secs, cur = defaultdict(list), None
        for l in CH.inputs(m)["diagram.txt"].splitlines():
            if l.startswith("UNIT "):
                cur = l[5:].strip()
            if cur:
                secs[cur].append(l.strip())
        diag[m] = secs

    def tokens(line):
        line = re.sub(r"([\w.()]+?)\.\{([^}]*)\}", lambda mm: " ".join(f"{mm.group(1)}.{f.strip()}" for f in mm.group(2).split(",")), line)
        return set(re.findall(r"[\w.]+", line))

    def capped(lines, n=6):
        return "\n".join(lines[:n]) + (f"\n(+{len(lines) - n} more lines)" if len(lines) > n else "")

    def input_facts(m, ent, el):
        x = names[m].get((ent, el)) or next((v for (e, n), v in names[m].items() if n == el), {}) or {}
        own = [rtl[m][l] for l in sorted(x.get("occurrence lines", [])) if l in rtl[m]]
        base, dl = el.split(".")[0], []
        for u, es in K.unit_entities(m).items():
            if ent in es:
                ports = {p for l in diag[m][u] if l.startswith(("INPUTS:", "OUTPUTS:")) for p in re.findall(r"[\w.]+", l.split(":", 1)[1])}
                dl += [l for l in diag[m][u] if el in tokens(l) or (base != el and base in ports and base in tokens(l))]
        kind = x.get("direction") or x.get("type") or ""
        return (f"{kind}; tags: {', '.join(x.get('role in RTL', [])) or '-'}\nRTL:\n{capped(own) or '-'}\n"
                f"Diagram:\n{capped(dl) or '-'}")

    credited = {(r["module"], r["matched reference entry"]): r["id"] for r in t["rows"] if r["verdict"] == "TP"}

    def reference_note(r):
        m, el = r["module"], r["element"]
        if r["verdict"] == "TP":
            return f"In the reference as '{r['matched reference entry']}'."
        base = el.split(".")[0]
        if any(rn == base and rn != el for rn, _ in refs.get(m, [])):
            return f"The reference lists the whole record '{base}' once; that one credit went to {credited.get((m, base), '?')}."
        return "Not in the reference list."

    def who_fp(rid):
        f = code[rid]["follows_prompt"]
        if f == "no":
            return "Model broke a rule"
        if f == "partly":
            return "Partly broke a rule"
        return "Reference choice" if groups[group_of[rid]]["cause"] == "reference convention" else "Prompt told it to"

    # ------------------------------------------------ rows: false positives, correct finds (and their evidence)
    row_ids, ev_rows, fp_rows, tp_rows = [], [], [], []
    for r in sorted(t["rows"], key=lambda r: (r["verdict"] != "FP", ORDER.index(group_of[r["id"]]) if r["id"] in group_of else 99,
                                              r["module"], r["element"])):
        m, ent, el = r["module"], r["entity"], r["element"]
        src = [a for a in recs[(m, 2)]["reply"]["structural assets"] if a.get("label") == "primary" and a.get("entity") == ent
               and (a.get("element") == el or el in (a.get("same role") or []))]
        pts = list(dict.fromkeys(p for a in src for p in (a.get("points") or []) if isinstance(p, str)))
        reasons = list(dict.fromkeys(a.get("reason", "") for a in src))
        concepts = "; ".join(f"{x['id']} {cname.get((m, x['id']), '')}" for x in r["concepts"])
        c = code.get(r["id"], {})
        ev_rows.append([r["id"], "false positive" if r["verdict"] == "FP" else "correct find", m.replace("neorv32_", ""), ent, el,
                        concepts, (f"({len(reasons)} reasons) " if len(reasons) > 1 else "") + "  ||  ".join(reasons),
                        ", ".join(map(str, r["phase 2 lines"])),
                        f"{len(pts)}: " + ", ".join(pts[:12]) + (f" (+{len(pts) - 12} more)" if len(pts) > 12 else ""),
                        input_facts(m, ent, el), reference_note(r), words(c.get("detail", ""))])
        row_ids.append(r["id"])
        kind = KIND.get(r["kind"], r["kind"]) + (" (whole record)" if r.get("is record") else "")
        if r["verdict"] == "FP":
            b = group_of[r["id"]]
            why = re.sub(r"^\s*(the\s+)?model(\s+says)?\s*:\s*", "", c.get("why_flagged", ""), flags=re.I)
            where = TEXT[b][3]
            if where.startswith("Prompt rule (code") and not removed_by[r["id"]] & set(RC.PROMPT_RULES):
                where = "Prompt rule (code test misses this row)"
            fp_rows.append([r["id"], GID[b], m.replace("neorv32_", ""), ent, el, kind, who_fp(r["id"]), words(why),
                            "Yes" if r["from 'same role' expansion"] else "", where,
                            ", ".join(sorted(removed_by[r["id"]])) or "No"])
        else:
            looks = [GID[b] for b in ORDER if r["id"] in (over[0][b] & over[1][b])]
            tp_rows.append([r["id"], m.replace("neorv32_", ""), ent, el, kind, concepts, "; ".join(looks)])

    # ------------------------------------------------ rows: missed entries (one precedence for category and text)
    def rel(a, b):
        return a == b or a.startswith(b + ".") or b.startswith(a + ".")
    fn_rows = []
    for k, f in enumerate(t["false negatives"]):
        m, rn = f["module"], f["reference entry"]
        p1, p2 = recs[(m, 1)]["reply"], recs[(m, 2)]["reply"]
        pat = re.compile(r"(?<![\w.])" + re.escape(rn) + r"(?![\w])")
        sa = [a for a in p2["structural assets"] if rel(a.get("element", ""), rn) or any(rel(x, rn) for x in (a.get("same role") or []))]
        pt_reason = [p for p in p2["points"] if pat.search(p.get("reason", ""))]
        pt_path = [p for p in p2["points"] if pat.search(p.get("reached by", ""))]
        merged = [x for x in p1.get("considered, not kept", []) if pat.search(json.dumps(x))]
        in_p1 = [c["id"] for c in p1["conceptual assets"] if pat.search(json.dumps(c))]
        x_in = CH.inputs(m)
        absent = not any(pat.search(x_in[q]) for q in x_in)
        parent = sorted({(e, n) for (e, n), v in names[m].items() for grp in (v.get("counterparts") or {}).values()
                         for q in grp if q.get("port") == rn and n in listed[(m, e)]})
        related = [r["id"] for r in t["rows"] if r["module"] == m and (r["entity"], r["element"]) in parent]
        judged = JUDGED_PAIRS.get((m, rn))
        if judged:
            related.append(judged[0])
        p1_note = ""
        if merged:
            x = merged[0]
            p1_note = f" In Phase 1 it was not a concept of its own: '{x.get('concept')}': {x.get('reason', '')}"
        if absent:
            stage, who, fix = "Not in this RTL", "Reference problem", "Correct the reference name."
            did = "The name occurs nowhere in the RTL, closed set or diagram." + (f" {judged[1]}" if judged else "")
        elif parent:
            stage, who, fix = "Listed under another name (this side of a port map)", "Naming (scoring)", f"Scoring alias ({GID['B08']})."
            did = f"Listed as {', '.join(n for _, n in parent)} ({', '.join(related)}): the same wire, named on this unit's side of the port map."
        elif sa and all(a["label"] != "primary" for a in sa):
            a = sa[0]
            stage, who = "Found, but labelled secondary", "Prompt rule (labelled secondary)"
            fix = "None drafted yet. Rules 2.6 and 2.7 push deciders, results and read-back words to secondary (see 'Prompt rules cited')."
            did = f"Labelled secondary for {a['concept id']} ({cname.get((m, a['concept id']), '')}): {a.get('reason', '')}"
        elif pt_reason:
            p = pt_reason[0]
            stage, who, fix = "Named in a Phase 2 point's reasoning, never listed", "Model", "None proposed yet."
            did = f"Named at point {p['point']} ({p['concept id']}): {p.get('reason', '')}" + p1_note
        elif pt_path:
            p = pt_path[0]
            ans = ", ".join(f"{q[0].upper()} {p.get(q)}" for q in ("confidentiality", "integrity", "availability"))
            stage, who, fix = "Named only in a Phase 2 walk path, judged not affected", "Model", "None proposed yet."
            did = f"Appears only in the path to point {p['point']} ({p['concept id']}): '{p.get('reached by', '').rstrip('.')}'. Answers there: {ans}." + p1_note
        elif merged or in_p1:
            stage, who, fix = "Left out in Phase 1", "Model", "None proposed yet."
            did = p1_note.strip() if merged else f"Mentioned only in the text of Phase 1 concept(s) {', '.join(in_p1)}."
        else:
            stage, who, fix, did = "Never mentioned by the model", "Model", "None proposed yet.", "-"
        ent = next((e for (e, n) in names[m] if n == rn), "")
        fid = f"FN{k:02d}"
        fn_rows.append([fid, m.replace("neorv32_", ""), rn, f["reference objective"], stage, who, words(did), ", ".join(related), fix])
        ev_rows.append([fid, "missed entry", m.replace("neorv32_", ""), ent, rn,
                        "; ".join(f"{c} {cname.get((m, c), '')}" for c in sorted({a['concept id'] for a in sa} | set(in_p1))),
                        "  ||  ".join(f"{a['label']} for {a['concept id']}: {a.get('reason', '')}" for a in sa), "", "",
                        input_facts(m, ent, rn) if ent else "-", f"In the reference (objective {f['reference objective']}).", words(did)])
        row_ids.append(fid)

    # ------------------------------------------------ rows: failure groups
    g_rows = []
    for b in ORDER:
        g = groups[b]
        mem = g["member_ids"]
        wc = Counter(who_fp(i) for i in mem)
        judge_rm = sum(1 for i in mem if any(i in supfp[x] for x in rule_of[b]))
        if TEXT[b][3] == "No fix":
            removes, hide = "-", "no fix"
        elif b == "B08":
            removes, hide = "0 (scoring change, see Proposed fix)", "-"
        else:
            own = CODE_TEST.get(b)
            removes = f"{sum(1 for i in mem if own in removed_by[i]) if own else '-'} / {judge_rm}"
            hide = "; ".join(f"{len(sup[0][x] & sup[1][x])} / {len(sup[0][x] | sup[1][x])}" for x in rule_of[b])
        g_rows.append([GID[b], TEXT[b][0], len(mem), wc.most_common(1)[0][0], "; ".join(f"{k} {v}" for k, v in wc.most_common()),
                       words(g["why_the_model_flags_these"]), TEXT[b][1], TEXT[b][2], TEXT[b][3], removes, hide,
                       f"{len(over[0][b] & over[1][b])} / {len(over[0][b] | over[1][b])}", b, ", ".join(g["example_ids"][:4])])

    # ------------------------------------------------ rows: calls
    warn_rows, call_rows = [], []
    for m in mods:
        for ph in (1, 2):
            r = recs[(m, ph)]
            u, rep = r["usage"], r["reply"] or {}
            if ph == 1:
                n = len(rep.get("conceptual assets", []))
                what = f"{n} concept{'s' if n != 1 else ''} kept, {len(rep.get('considered, not kept', []))} left out"
            else:
                what = f"{len(rep.get('points', []))} points, {sum(a.get('label') == 'primary' for a in rep.get('structural assets', []))} primary rows"
            call_rows.append([m.replace("neorv32_", ""), ph, what, score["per_module"][m]["emit"] if ph == 2 else "", len(r["problems"]),
                              u["in"], u.get("cached") or 0, u["out"], CE.usd(u)])
            warn_rows += [[m.replace("neorv32_", ""), ph, w] for w in r["problems"]]

    # ------------------------------------------------ workbook
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HEAD = PatternFill("solid", fgColor="1F3864")
    LINK = Font(color="0563C1", underline="single")
    thin = Side(style="thin", color="D9D9D9")
    wb = Workbook()
    ev_row_of = {rid: i + 2 for i, rid in enumerate(row_ids)}

    def table(ws, header, rows_, widths, color_cols=(), link_col=None, wrap=True):
        for j, h in enumerate(header, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(rows_, 2):
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.alignment = Alignment(wrap_text=wrap, vertical="top")
                c.border = Border(bottom=thin)
            for col in color_cols:
                v = str(row[col - 1])
                fill = WHO_FILL.get(v) or next((FIX_FILL[k] for k in FIX_FILL if v.startswith(k)), None)
                if fill:
                    ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=fill)
            if link_col and row[link_col - 1] in ev_row_of:
                c = ws.cell(row=i, column=link_col)
                c.hyperlink = f"#'Evidence'!A{ev_row_of[row[link_col - 1]]}"
                c.font = LINK
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{1 + len(rows_)}"

    # Start here
    ws = wb.active
    ws.title = "Start here"
    fps_ids = [r["id"] for r in t["rows"] if r["verdict"] == "FP"]
    who_c = Counter(who_fp(i) for i in fps_ids)
    partly = [i for i in fps_ids if who_fp(i) == "Partly broke a rule"]
    partly_broke = sum(1 for i in partly if re.search(r"\bbr(oke|eaks?)\b", code[i]["prompt_rule"], re.I))
    ref_ids = [i for i in fps_ids if who_fp(i) == "Reference choice"]
    ref_names = sum(1 for i in ref_ids if group_of[i] == "B08")
    ref_unit = sum(1 for i in ref_ids if group_of[i] == "B09")
    per_mod = Counter(rows[i]["module"].replace("neorv32_", "") for i in fps_ids)
    bus_copies = sum(1 for i in fps_ids if rows[i]["module"] == "neorv32_bus" and rows[i]["from 'same role' expansion"])
    look_any = sum(1 for r in t["rows"] if r["verdict"] == "TP" and any(r["id"] in (over[0][b] & over[1][b]) for b in groups))
    absent_n = sum(1 for r in fn_rows if r[4] == "Not in this RTL")
    findable = score["ref"] - absent_n
    fn_c = Counter(r[4] for r in fn_rows)
    fam = Counter(fix_family(TEXT[group_of[i]][3]) for i in fps_ids)
    gfam = Counter(fix_family(TEXT[b][3]) for b in ORDER)
    (tn, tu, th), (hn, hu, hh) = exposure["tuning"], exposure["heldout"]
    lines = [
        ("Prompt failure analysis", "title"),
        ("What this run did", "h"),
        ("The prompt asks a model to list, for each RTL module, the elements that hold or decide security-relevant values (the 'primary assets').", ""),
        (f"Model {model_name}; prompt two_phase_v1, two calls per module (Phase 1 names what is worth protecting, Phase 2 finds where it "
         f"is in the RTL); the 15 tuning modules; run on 2026-10-11.", ""),
        ("The list is scored against the project's reference list (the ground truth). Terms are explained on the Glossary sheet.", ""),
        ("", ""),
        ("Result", "h"),
        (f"Listed {score['emit']} elements: {score['tp']} correct finds and {score['fp']} false positives.", ""),
        (f"Precision {score['precision']:.3f} ({score['tp']} of {score['emit']} listed are correct).", ""),
        (f"Recall {score['recall']:.3f} ({score['tp']} of the {score['ref']} reference entries found). {absent_n} entries are names that occur "
         f"nowhere in this RTL; of the {findable} that could be found, {score['tp'] / findable:.3f}.", ""),
        (f"Cost USD {total_usd:.2f} for 30 calls.", ""),
        ("False positives per module: " + ", ".join(f"{k} {v}" for k, v in per_mod.most_common()) +
         f". In bus, {bus_copies} of the {per_mod['bus']} are 'same role' copies: one listing repeated across identical ports.", ""),
        ("", ""),
        ("Why the false positives happened", "h"),
        ("AI agent readers read each false positive together with the model's own reason (not a human audit; method in ERROR_ANALYSIS.md).", ""),
        (f"Prompt told it to: {who_c['Prompt told it to']}. The model followed the prompt; the prompt's own definitions make these primary.", "Prompt told it to"),
        (f"Reference choice: {who_c['Reference choice']}. The model followed the prompt, but the reference keeps some elements of this shape "
         f"and not others ({len(ref_ids) - ref_names - ref_unit}), names the same thing differently ({ref_names}), or leaves the unit out "
         f"({ref_unit}).", "Reference choice"),
        (f"Partly broke a rule: {who_c['Partly broke a rule']}. The model followed one clause and broke another ({partly_broke}), or two "
         f"clauses point different ways ({len(partly) - partly_broke}).", "Partly broke a rule"),
        (f"Model broke a rule: {who_c['Model broke a rule']}. The prompt says not to list them.", "Model broke a rule"),
        (f"{look_any} of the {score['tp']} correct finds look like at least one failure group (both judges agree). This is why some groups "
         f"have no safe fix.", ""),
        ("", ""),
        ("What the fixes would do", "h"),
        (f"Prompt rules ({gfam['Prompt rule']} groups) cover {fam['Prompt rule']} false positives; scoring changes "
         f"({gfam['Scoring filter'] + gfam['Scoring alias']} groups) cover {fam['Scoring filter'] + fam['Scoring alias']}; no fix "
         f"({gfam['No fix']} groups) for {fam['No fix']}.", ""),
        (f"Estimate from filtering the stored output with the code tests (in-sample, one run, not a re-run): the prompt rules alone give "
         f"precision {ph_prompt['precision']:.3f} (recall {ph_prompt['recall']:.3f}); with the scoring filters, precision "
         f"{ph_all['precision']:.3f} (recall {ph_all['recall']:.3f}).", ""),
        (f"Risk to correct finds: the code tests remove none of this run's {score['tp']} correct finds. Tested on every reference entry as if "
         f"it were listed, they would hide {th} of the {tn - tu} tuning entries and {hh} of the {hn - hu} held-out entries ({tu} and {hu} "
         f"reference names are not in the closed sets). Two judges also read every correct find against the drafted rules' wording: when "
         f"both agree, none is hidden; one judge alone flagged 3 borderline cases.", ""),
        ("Only a re-run measures the real effect; one re-run costs about USD 27. No fix is drafted yet for the missed entries labelled secondary.", ""),
        ("", ""),
        ("Why entries were missed", "h"),
    ] + [(f"{k}: {v}", "") for k, v in fn_c.most_common()] + [
        ("", ""),
        ("Sheets, in reading order (click to open)", "h"),
        ("Glossary: the terms used.", "link:Glossary"),
        (f"Failure groups: the false positives in {len(ORDER)} groups (G01 to G{len(ORDER):02d}), with cause, fix and the fix's effect.", "link:Failure groups"),
        ("False positives: one row per element. Click an Id to open its evidence.", "link:False positives"),
        ("Missed entries: one row per reference entry the model did not list.", "link:Missed entries"),
        ("Correct finds: the elements the model got right, and the failure groups each one resembles.", "link:Correct finds"),
        ("Evidence: the full detail of every case (false positives, correct finds, missed entries).", "link:Evidence"),
        ("Prompt rules cited: the prompt clauses this analysis refers to, quoted word for word.", "link:Prompt rules cited"),
        ("Run setup: model, prompt, inputs and settings, in the MODEL / PROMPT / INPUT / CONTEXT / TOOLS / RUNTIME / EXPECTED / OBSERVED form.", "link:Run setup"),
        ("Calls: tokens, cost and check warnings for each of the 30 calls.", "link:Calls"),
        ("The exact prompts, inputs and replies are in the companion file PROMPT_FAILURE_EXACT_TEXT.xlsx.", ""),
        ("", ""),
        ("Colours", "h"),
        ("Who: Prompt told it to", "Prompt told it to"), ("Who: Partly broke a rule", "Partly broke a rule"),
        ("Who: Model broke a rule", "Model broke a rule"), ("Who: Reference choice", "Reference choice"),
        ("Fix: prompt rule", "fix:Prompt rule"), ("Fix: scoring filter or alias", "fix:Scoring filter"), ("Fix: none", "fix:No fix"),
    ]
    for i, (txt, style) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            c.font = Font(bold=True, size=15)
        elif style == "h":
            c.font = Font(bold=True, size=12, color="1F3864")
        elif style in WHO_FILL:
            c.fill = PatternFill("solid", fgColor=WHO_FILL[style])
        elif style.startswith("fix:"):
            c.fill = PatternFill("solid", fgColor=FIX_FILL[style[4:]])
        elif style.startswith("link:"):
            c.hyperlink = f"#'{style[5:]}'!A1"
            c.font = LINK
    ws.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Glossary")
    table(ws, ["Term", "Meaning"], GLOSSARY, [26, 120])
    ws = wb.create_sheet("Failure groups")
    table(ws, ["Group", "What it is", "False positives", "Mostly who", "Who, case by case", "Why the model listed them (AI reader summary)",
               "What should happen", "Proposed fix", "Fix belongs in",
               "False positives the fix removes in this run (code test / one judge reading the wording)",
               "Correct finds the fix would hide (two judges: both agree / at least one)",
               "Correct finds that look like this group (both judges / at least one; one find can match several groups)",
               "Old code", "Example ids"],
          g_rows, [7, 32, 9, 15, 30, 70, 46, 46, 18, 16, 16, 18, 7, 22], color_cols=(4, 9))
    ws = wb.create_sheet("False positives")
    table(ws, ["Id", "Group", "Module", "Entity (sub-unit)", "Element", "Kind", "Who",
               "Why it was listed (AI reader's summary of the model's reason; the model's own words are in Evidence)",
               "'Same role' copy", "Fix belongs in", "Removed by a code test? (which)"],
          fp_rows, [7, 7, 12, 22, 24, 18, 18, 70, 9, 22, 14], color_cols=(7, 10), link_col=1)
    ws = wb.create_sheet("Missed entries")
    table(ws, ["Id", "Module", "Reference entry", "Reference objective", "Where it was lost", "Who", "What the model did instead",
               "Related false positives (same wire or same thing)", "Possible fix"],
          fn_rows, [7, 12, 20, 12, 30, 20, 80, 14, 40], color_cols=(6,), link_col=1)
    ws = wb.create_sheet("Correct finds")
    table(ws, ["Id", "Module", "Entity (sub-unit)", "Element", "Kind", "Concept(s) it was listed for",
               "Looks like failure group(s) (both judges agree)"], tp_rows, [7, 12, 22, 24, 18, 50, 18], link_col=1)
    ws = wb.create_sheet("Evidence")
    table(ws, ["Id", "Case", "Module", "Entity", "Element", "Concept(s)", "Model's reason(s) (several are separated by ||)",
               "RTL lines the model cited", "Points the model reached (count: names)",
               "Input facts (closed-set kind and tags, RTL lines, diagram lines)", "Reference", "AI reader's note"],
          ev_rows, [7, 12, 11, 18, 20, 28, 70, 14, 30, 60, 26, 60])
    ws = wb.create_sheet("Prompt rules cited")
    crow = []
    for name, fn, a, z, plain in CLAUSES:
        txt = (PROMPTS / fn).read_text(encoding="utf-8")
        ma = re.search(a, txt, re.M)
        assert ma, f"clause not found: {name}"
        mz = re.search(z, txt[ma.end():], re.M)
        crow.append([name, plain, fn, txt[ma.start(): ma.end() + (mz.start() if mz else len(txt))].strip()])
    table(ws, ["Clause", "In plain words", "File", "Exact text"], crow, [30, 50, 12, 110])
    ws = wb.create_sheet("Run setup")
    setup = [
        ("MODEL", f"{model_name} (the exact version the API returned, the same for all 30 calls)."),
        ("PROMPT", "two_phase_v1. Call 1 (Phase 1): common.md + phase1.md. Call 2 (Phase 2): common.md + phase2.md. Folder "
                   "experiments/03_two_phase_csa_pio/lab/prompts/two_phase_v1. Exact text in the companion file."),
        ("INPUT", "Per module: the RTL (comments removed, lines numbered), the closed set (JSON) and the block diagram (text); inputs "
                  "version v1 (experiments/03_two_phase_csa_pio/lab/inputs/v1). Exact user messages in the companion file."),
        ("CONTEXT", "No retrieved documents, no chat history, no memory, no examples; each call is a single turn. Phase 2 also receives "
                    "Phase 1's list of concepts inside its user message."),
        ("TOOLS", "None. Output format: JSON mode (any JSON object; no schema)."),
        ("RUNTIME", "Reasoning effort high. Output limit 128,000 tokens. Background mode (the request runs on the server and is polled). "
                    "No temperature or truncation setting sent. All 30 calls finished with status 'completed'."),
        ("EXPECTED", "The reference list for each module, scored with the strict scorer (src/eval_assets.py): a listed element is correct "
                     "when it matches a reference entry by name; a whole record also matches one of its fields, and each entry gives one credit."),
        ("OBSERVED", f"{score['emit']} elements listed: {score['tp']} correct, {score['fp']} false positives, {score['fn']} missed. "
                     f"Precision {score['precision']:.3f}, recall {score['recall']:.3f}. Cost USD {total_usd:.2f}."),
    ]
    table(ws, ["Field", "This run"], setup, [14, 130])
    for i in range(2, len(setup) + 2):
        ws.cell(row=i, column=1).font = Font(bold=True)
    ws = wb.create_sheet("Calls")
    tot = ["TOTAL", "", "", sum(r[3] for r in call_rows if r[3] != ""), sum(r[4] for r in call_rows), sum(r[5] for r in call_rows),
           sum(r[6] for r in call_rows), sum(r[7] for r in call_rows), round(total_usd, 2)]
    table(ws, ["Module", "Phase", "What it returned", "Elements scored (after 'same role' copies are added and repeats merged)",
               "Check warnings", "Input tokens", "Cached input tokens", "Output tokens", "USD"],
          [r[:8] + [round(r[8], 4)] for r in call_rows] + [tot], [14, 7, 34, 16, 10, 12, 12, 12, 9], wrap=False)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.auto_filter.ref = f"A1:I{1 + len(call_rows)}"
    r0 = ws.max_row + 2
    ws.cell(row=r0, column=1, value="Check warnings (from the runner's checks of each reply)").font = Font(bold=True)
    for j, w in enumerate(warn_rows, r0 + 1):
        for k_, v in enumerate(w, 1):
            ws.cell(row=j, column=k_, value=v)
    dest = HERE / "PROMPT_FAILURE_DIAGNOSIS.xlsx"
    wb.save(dest)

    # ------------------------------------------------ companion file: one line per row, readable sheet names, links
    wx = Workbook()
    s0 = wx.active
    s0.title = "Contents"
    for j, h in enumerate(["Sheet", "Text", "Lines", "Characters", "Fingerprint (first 12 hex of SHA-256; equal = identical texts)"], 1):
        c = s0.cell(row=1, column=j, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for i, (sheet_name, desc, txt, sh) in enumerate(exact, 2):
        s = wx.create_sheet(sheet_name)
        s["A1"] = (f"{desc}: {len(txt):,} characters, fingerprint {sh}. One line of the text per row; a line longer than "
                   f"30,000 characters (an Excel cell limit) continues on the next row(s), marked 'continued' in column B.")
        s["A1"].font = Font(bold=True)
        r = 2
        for line in txt.split("\n"):
            for q, p in enumerate([line[q:q + 30000] for q in range(0, len(line), 30000)] or [""]):
                s.cell(row=r, column=1, value=p)
                if q:
                    s.cell(row=r, column=2, value="continued")
                r += 1
        s.column_dimensions["A"].width = 160
        s0.cell(row=i, column=1, value=sheet_name).hyperlink = f"#'{sheet_name}'!A1"
        s0.cell(row=i, column=1).font = LINK
        for j, v in enumerate([desc, txt.count("\n") + 1, len(txt), sh], 2):
            s0.cell(row=i, column=j, value=v)
    for j, w in enumerate([22, 46, 10, 12, 34], 1):
        s0.column_dimensions[get_column_letter(j)].width = w
    s0.auto_filter.ref = f"A1:E{len(exact) + 1}"
    s0.freeze_panes = "A2"
    wx.save(HERE / "PROMPT_FAILURE_EXACT_TEXT.xlsx")
    print(f"{dest}\n{HERE / 'PROMPT_FAILURE_EXACT_TEXT.xlsx'} ({len(exact)} texts)")


if __name__ == "__main__":
    main()
